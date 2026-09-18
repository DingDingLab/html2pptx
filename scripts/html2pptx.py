#!/usr/bin/env python
"""HTML 幻灯片 → 图片版 PPTX（排版 100% 保真，文字不可再编辑，适合交作业/交付）。

用法:
  python html2pptx.py input.html [output.pptx]

依赖: playwright + python-pptx + Pillow
前置: 脚本假设 HTML 使用 templates/presentation.html 的结构（.slide 元素 + active 类切换）
"""
import sys, pathlib
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Inches

html_path = sys.argv[1]
out_path = sys.argv[2] if len(sys.argv) > 2 else str(pathlib.Path(html_path).with_suffix(".pptx"))

# 16:9 逻辑尺寸（对应 HTML 的 960x540 视觉比例）
SLIDE_W, SLIDE_H = Inches(10), Inches(5.625)
VIEW_W, VIEW_H = 1280, 720

with sync_playwright() as p:
    b = p.chromium.launch()
    pg = b.new_page(viewport={"width": VIEW_W, "height": VIEW_H})
    pg.goto(f"file://{pathlib.Path(html_path).resolve()}")
    pg.wait_for_timeout(600)

    n = pg.evaluate("document.querySelectorAll('.slide').length")
    print(f"检测到 {n} 页")

    shots = []
    for i in range(n):
        # 先移除所有 active，再只激活第 i 页，等过渡结束再截（否则上一页会残留叠加）
        pg.evaluate("document.querySelectorAll('.slide').forEach(s => s.classList.remove('active'))")
        pg.evaluate(f"document.querySelectorAll('.slide')[{i}].classList.add('active')")
        pg.wait_for_timeout(650)
        tmp = f"/tmp/ppt_slide_{i:02d}.png"
        pg.screenshot(path=tmp)
        shots.append(tmp)
        print(f"  slide {i+1} -> {tmp}")
    b.close()

prs = Presentation()
prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
blank = prs.slide_layouts[6]
for s in shots:
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(s, 0, 0, width=SLIDE_W, height=SLIDE_H)

prs.save(out_path)
print(f"已生成 {out_path}（{len(shots)} 页 PPTX）")
